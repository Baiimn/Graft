#!/usr/bin/env python3
"""Convert a repo's documents into shadow markdown that graft's prose tier indexes.

Each .pdf / .docx / .pptx under <root> becomes <root>/graft/shadow/<same relative
path>.md. graft then indexes the shadow like any other markdown file, so
`graft ask` can answer from a report, a deck or a lecture PDF. The shadow is
text only: figures, layout and fonts still need the real file.

Run on demand -- this is deliberately NOT wired into `graft build`, which would
make every repo depend on Python.

    python scripts/graft-docs.py D:/Vault
    python scripts/graft-docs.py --self-check
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

SHADOW = "graft/shadow"
MANIFEST = ".manifest.json"
SKIP_DIRS = {"node_modules", ".git", "__pycache__", ".tmp", "graft"}
EXTS = {".pdf", ".docx", ".pptx"}


def sources(root: Path) -> list[Path]:
    out: list[Path] = []

    def walk(d: Path) -> None:
        for e in sorted(d.iterdir()):
            if e.is_dir():
                # Hidden dirs are tooling scratch (.git, .playwright-mcp, .venv):
                # never his documents, always junk in the index.
                if e.name not in SKIP_DIRS and not e.name.startswith("."):
                    walk(e)
            elif e.suffix.lower() in EXTS and not e.name.startswith("~$"):
                out.append(e)

    walk(root)
    return out


def digest(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def md_table(rows: list[list[str]]) -> list[str]:
    """Rows -> a markdown table. First row is the header."""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    cells = [[c.replace("|", "\\|").replace("\n", " ").strip() for c in r] for r in rows]
    out = ["| " + " | ".join(cells[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in cells[1:]]
    return out + [""]


def convert_pdf(p: Path) -> tuple[list[str], int]:
    """Body lines and page count. A page with no text layer is skipped."""
    import fitz

    body: list[str] = []
    with fitz.open(p) as doc:
        pages = doc.page_count
        for i, page in enumerate(doc, 1):
            text = page.get_text("text").strip()
            if not text:
                continue
            body += [f"## p.{i}", "", text, ""]
    return body, pages


def convert_docx(p: Path) -> tuple[list[str], int]:
    import docx

    d = docx.Document(str(p))
    body: list[str] = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(block, d)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style.startswith("Heading"):
                tail = style.split()[-1]
                level = int(tail) if tail.isdigit() else 1
                body += ["#" * min(level, 6) + f" {text}", ""]
            else:
                body += [text, ""]
        elif tag == "tbl":
            table = docx.table.Table(block, d)
            body += md_table([[c.text for c in row.cells] for row in table.rows])
    return body, len(d.paragraphs)


def convert_pptx(p: Path) -> tuple[list[str], int]:
    from pptx import Presentation

    prs = Presentation(str(p))
    body: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip().replace("\n", " ")
        body += [f"## Slide {i}" + (f" - {title}" if title else ""), ""]
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False):
                body += md_table([[c.text for c in row.cells] for row in shape.table.rows])
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    body += [text, ""]
    return body, len(prs.slides)


CONVERTERS = {".pdf": convert_pdf, ".docx": convert_docx, ".pptx": convert_pptx}
UNITS = {".pdf": "pages", ".docx": "paragraphs", ".pptx": "slides"}


def convert(src: Path, rel: str) -> str | None:
    """Shadow markdown for one document, or None if it carries no text."""
    ext = src.suffix.lower()
    body, count = CONVERTERS[ext](src)
    if not body:
        return None
    header = [f"# {src.stem}", "", f"> source: {rel} - {count} {UNITS[ext]}", ""]
    return "\n".join(header + body).rstrip() + "\n"


def run(root: Path) -> int:
    shadow_root = root / SHADOW
    manifest_path = shadow_root / MANIFEST
    manifest: dict[str, str] = {}
    if manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text("utf8"))
        except json.JSONDecodeError:
            manifest = {}

    found = sources(root)
    seen: dict[str, str] = {}
    written = skipped = 0
    unreadable: list[str] = []

    for src in found:
        rel = src.relative_to(root).as_posix()
        out = shadow_root / (rel + ".md")
        h = digest(src)
        seen[rel] = h
        if manifest.get(rel) == h and out.exists():
            skipped += 1
            continue
        try:
            text = convert(src, rel)
        except Exception as exc:  # a broken document must not kill the run
            unreadable.append(f"{rel} - {type(exc).__name__}: {exc}")
            seen.pop(rel)
            continue
        if text is None:
            unreadable.append(f"{rel} - no text layer (scanned?)")
            seen.pop(rel)
            continue
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf8")
        written += 1

    # A shadow whose source is gone answers with text that no longer exists.
    pruned = 0
    if shadow_root.exists():
        live = {shadow_root / (rel + ".md") for rel in seen}
        for f in shadow_root.rglob("*.md"):
            if f not in live:
                f.unlink()
                pruned += 1
        for d in sorted(shadow_root.rglob("*"), key=lambda x: -len(x.parts)):
            if d.is_dir() and not any(d.iterdir()):
                d.rmdir()

    shadow_root.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(seen, indent=2, sort_keys=True), encoding="utf8")

    print(f"{len(found)} documents | {written} written | {skipped} unchanged | {pruned} pruned")
    if unreadable:
        print(f"\ncould not read ({len(unreadable)}):")
        for line in unreadable:
            print(f"  {line}")
    print(f"\nnow run: graft build   (in {root})")
    return 0


def self_check() -> int:
    """Build one tiny doc of each kind, convert it, assert the text survives."""
    import tempfile

    import docx as docx_mod
    import fitz
    from pptx import Presentation
    from pptx.util import Inches

    sentence = "Hypothesis testing rejects the null at alpha 0.05."
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        d = docx_mod.Document()
        d.add_heading("Chapter Five", level=2)
        d.add_paragraph(sentence)
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "stat"
        t.cell(0, 1).text = "value"
        t.cell(1, 0).text = "p"
        t.cell(1, 1).text = "0.03"
        d.save(root / "a.docx")

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Findings"
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = sentence
        prs.save(root / "b.pptx")

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), sentence)
        doc.save(root / "c.pdf")
        doc.close()

        run(root)
        shadow = root / SHADOW

        a = (shadow / "a.docx.md").read_text("utf8")
        assert "## Chapter Five" in a, a
        assert sentence in a, a
        assert "| stat | value |" in a, a

        b = (shadow / "b.pptx.md").read_text("utf8")
        assert "## Slide 1 - Findings" in b, b
        assert sentence in b, b

        c = (shadow / "c.pdf.md").read_text("utf8")
        assert "## p.1" in c, c
        assert sentence in c, c
        assert "> source: c.pdf - 1 pages" in c, c

        # incremental: a second run rewrites nothing
        before = {f: f.stat().st_mtime_ns for f in shadow.rglob("*.md")}
        run(root)
        assert {f: f.stat().st_mtime_ns for f in shadow.rglob("*.md")} == before, "second run rewrote files"

        # orphan: source deleted -> shadow deleted
        (root / "a.docx").unlink()
        run(root)
        assert not (shadow / "a.docx.md").exists(), "orphan shadow survived"

    print("self-check OK")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("root", nargs="?", default=".", help="repo root to convert")
    ap.add_argument("--self-check", action="store_true", help="run the built-in check and exit")
    args = ap.parse_args()
    if args.self_check:
        return self_check()
    root = Path(args.root).resolve()
    if not root.is_dir():
        print(f"not a directory: {root}", file=sys.stderr)
        return 1
    return run(root)


if __name__ == "__main__":
    raise SystemExit(main())
